"""Unit tests for ingestion/attendance.py.

Uses a tiny synthetic SAR workbook (built in tests/conftest.py via openpyxl —
no real Premier DLC data or PII involved) to exercise the layout-tolerant
parser, the anonymisation default, the missing-school handling, and the
aggregate statistics (annual mean, pp-change, lost child-school-days).
"""
from __future__ import annotations

import datetime as dt

import openpyxl
import pytest

from ingestion import attendance as att


@pytest.fixture(autouse=True)
def _patch_roster_csv(monkeypatch, mini_schools_csv):
    """Point the module's hardcoded roster path at our synthetic 3-school CSV
    for every test in this file, instead of the real 50-school roster."""
    monkeypatch.setattr(att, "SCHOOLS_CSV", mini_schools_csv)


class TestLoadRoster:
    def test_loads_synthetic_roster(self):
        roster = att._load_roster()
        assert set(roster.keys()) == {90000001, 90000002, 90000003}
        assert roster[90000001]["students"] == 100
        assert roster[90000001]["cluster"] == "C-1"


class TestSha256:
    def test_deterministic_for_same_file(self, mini_sar_xlsx):
        h1 = att._sha256(mini_sar_xlsx)
        h2 = att._sha256(mini_sar_xlsx)
        assert h1 == h2
        assert len(h1) == 64  # hex sha256

    def test_differs_for_different_content(self, mini_sar_xlsx, tmp_path):
        other = tmp_path / "other.xlsx"
        other.write_bytes(mini_sar_xlsx.read_bytes() + b"\x00")
        assert att._sha256(mini_sar_xlsx) != att._sha256(other)


class TestFindHeaderRow:
    def test_locates_emis_cell_case_insensitively(self):
        rows = [
            ("title", None),
            ("emis", "School Name"),
            (90000001, "Test A"),
        ]
        ri, ci = att._find_header_row(rows)
        assert (ri, ci) == (1, 0)

    def test_raises_when_no_emis_column_present(self):
        rows = [("title",), ("Name", "Value")] + [("x",)] * 14
        with pytest.raises(ValueError, match="Could not find EMIS column"):
            att._find_header_row(rows)


class TestParseYearSheet:
    def test_raises_without_datetime_month_columns(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="EMIS")
        ws.cell(row=1, column=2, value="Jan")  # a *string*, not a datetime cell
        ws.cell(row=2, column=1, value=90000001)
        ws.cell(row=2, column=2, value=0.9)
        with pytest.raises(ValueError, match="No month columns"):
            att._parse_year_sheet(ws)

    def test_ignores_out_of_range_fractions(self):
        """Only numeric cells within [0, 1] are treated as attendance fractions
        — this guards against stray non-attendance numbers (e.g. a totals row
        or an accidentally-typed percentage like 90 instead of 0.90)."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="EMIS")
        ws.cell(row=1, column=2, value=dt.datetime(2023, 1, 1))
        ws.cell(row=2, column=1, value=90000001)
        ws.cell(row=2, column=2, value=90)  # out of range — should be dropped
        months, out = att._parse_year_sheet(ws)
        assert months == ["Jan"]
        assert out.get(90000001, {}) == {}


class TestFindYearSheet:
    def test_exact_match(self, mini_sar_xlsx):
        wb = openpyxl.load_workbook(mini_sar_xlsx, read_only=True, data_only=True)
        ws = att._find_year_sheet(wb, "2023")
        assert ws.title == "2023"

    def test_fuzzy_match_when_no_exact_sheet(self):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        wb.create_sheet("SAR 2023 (final)")
        ws = att._find_year_sheet(wb, "2023")
        assert ws.title == "SAR 2023 (final)"

    def test_raises_keyerror_when_year_absent(self):
        wb = openpyxl.Workbook()
        with pytest.raises(KeyError):
            att._find_year_sheet(wb, "2099")


class TestBuildPayloadAnonymisation:
    def test_default_anonymises_school_ids(self, mini_sar_xlsx):
        payload = att.build_payload(mini_sar_xlsx)
        assert payload["meta"]["school_id_anonymised"] is True
        assert set(payload["schools"].keys()) == {"school_id_01", "school_id_02", "school_id_03"}
        # Anonymised IDs are assigned in stable EMIS-sorted order.
        assert payload["meta"]["emis_to_id"]["90000001"] == "school_id_01"
        assert payload["meta"]["emis_to_id"]["90000003"] == "school_id_03"

    def test_reveal_emis_uses_real_codes_as_keys(self, mini_sar_xlsx):
        payload = att.build_payload(mini_sar_xlsx, reveal_emis=True)
        assert payload["meta"]["school_id_anonymised"] is False
        assert set(payload["schools"].keys()) == {"90000001", "90000002", "90000003"}


class TestBuildPayloadContent:
    @pytest.fixture
    def payload(self, mini_sar_xlsx):
        return att.build_payload(mini_sar_xlsx)

    def test_source_hash_recorded_for_reproducibility(self, payload, mini_sar_xlsx):
        assert payload["meta"]["source_sha256"] == att._sha256(mini_sar_xlsx)

    def test_school_c_missing_from_2025_is_recorded(self, payload):
        # School C (school_id_03, emis 90000003) exists in 2023/2024 but was
        # dropped from the synthetic 2025 sheet.
        assert "school_id_03" in payload["meta"]["missing_school_ids"]["2025"]
        rec = payload["schools"]["school_id_03"]
        assert rec["monthly"]["2025"] is None
        assert rec["annual_mean"]["2025"] is None
        # No 2025 mean -> the 2023->2025 delta can't be computed for this school.
        assert rec["pp_change_2023_2025"] is None

    def test_present_schools_have_correct_annual_means(self, payload):
        # School A 2023: Jan 0.90, Feb 0.92 -> mean 0.91
        rec_a = payload["schools"]["school_id_01"]
        assert rec_a["annual_mean"]["2023"] == pytest.approx(0.91, abs=1e-4)
        # School B 2023: Jan 0.80, Feb 0.82 -> mean 0.81
        rec_b = payload["schools"]["school_id_02"]
        assert rec_b["annual_mean"]["2023"] == pytest.approx(0.81, abs=1e-4)

    def test_pp_change_matches_hand_computed_value(self, payload):
        # School A: 2023 mean 0.91, 2025 mean (0.95+0.90)/2 = 0.925
        # pp change = (0.925 - 0.91) * 100 = 1.5
        rec_a = payload["schools"]["school_id_01"]
        assert rec_a["pp_change_2023_2025"] == pytest.approx(1.5, abs=0.05)

    def test_schools_covered_counts_only_non_null_years(self, payload):
        assert payload["meta"]["schools_covered"]["2023"] == 3
        assert payload["meta"]["schools_covered"]["2024"] == 3
        assert payload["meta"]["schools_covered"]["2025"] == 2  # school C dropped

    def test_lost_child_school_days_uses_common_months_only(self, payload):
        # Only schools A & B have both 2023 and 2025 data (school C is missing
        # in 2025). Common months for both are Jan & Feb.
        # A: students=100; delta Jan = 0.90-0.95=-0.05; delta Feb=0.92-0.90=0.02
        # B: students=200; delta Jan = 0.80-0.60=0.20; delta Feb=0.82-0.65=0.17
        expected = (
            100 * -0.05 * att.SCHOOL_DAYS_PER_MONTH
            + 100 * 0.02 * att.SCHOOL_DAYS_PER_MONTH
            + 200 * 0.20 * att.SCHOOL_DAYS_PER_MONTH
            + 200 * 0.17 * att.SCHOOL_DAYS_PER_MONTH
        )
        assert payload["aggregates"]["lost_child_school_days_2023_baseline"] == round(expected)

    def test_by_cluster_aggregates_only_include_this_clusters_schools(self, payload):
        # Schools A & B are cluster C-1, school C is cluster C-2 (but has no
        # 2025 data, so by_cluster["C-2"] must not have a 2025 key).
        assert "C-1" in payload["aggregates"]["by_cluster"]
        assert "C-2" in payload["aggregates"]["by_cluster"]
        assert "2025" not in payload["aggregates"]["by_cluster"]["C-2"]
        assert "2025" in payload["aggregates"]["by_cluster"]["C-1"]

    def test_methodology_and_provider_metadata_present(self, payload):
        assert payload["meta"]["provider"] == "Premier DLC"
        assert payload["meta"]["methodology_url"] == "docs/methodology-attendance.md"
        assert payload["meta"]["school_days_per_month"] == att.SCHOOL_DAYS_PER_MONTH


class TestMainCLI:
    def test_returns_2_when_xlsx_missing(self, monkeypatch, tmp_path):
        monkeypatch.setattr(
            "sys.argv", ["attendance.py", "--xlsx", str(tmp_path / "nope.xlsx")]
        )
        assert att.main() == 2

    def test_writes_output_file_end_to_end(self, monkeypatch, mini_sar_xlsx, tmp_path):
        out_path = tmp_path / "out" / "attendance.json"
        monkeypatch.setattr(
            "sys.argv",
            ["attendance.py", "--xlsx", str(mini_sar_xlsx), "--out", str(out_path)],
        )
        assert att.main() == 0
        assert out_path.exists()
        import json
        payload = json.loads(out_path.read_text())
        assert payload["meta"]["schools_covered"]["2023"] == 3
