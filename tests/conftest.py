"""Shared fixtures for the school-climate-hub test suite.

No real PII or production data is used here. School names, EMIS codes, and
attendance numbers below are synthetic and only shaped to match the formats
that `scoring/vulnerability.py` and `ingestion/attendance.py` expect.
"""
from __future__ import annotations

import csv
import datetime as dt
from pathlib import Path

import openpyxl
import pytest


@pytest.fixture
def mini_schools() -> list[dict]:
    """A tiny synthetic 3-school roster shaped like scoring.vulnerability.load_schools()."""
    return [
        {
            "emis_code": 90000001,
            "school_name": "Test Primary A",
            "project": "PSSP",
            "cluster": "C-1",
            "tehsil": "Test Tehsil",
            "district": "Test District",
            "students": 100,
            "lat": 32.08,
            "lon": 74.15,
        },
        {
            "emis_code": 90000002,
            "school_name": "Test Primary B",
            "project": "PSSP",
            "cluster": "C-1",
            "tehsil": "Test Tehsil",
            "district": "Test District",
            "students": 200,
            "lat": 32.10,
            "lon": 74.18,
        },
        {
            "emis_code": 90000003,
            "school_name": "Test Primary C",
            "project": "PSRP",
            "cluster": "C-2",
            "tehsil": "Test Tehsil",
            "district": "Test District",
            "students": 50,
            "lat": 32.20,
            "lon": 74.30,
        },
    ]


@pytest.fixture
def mini_schools_csv(tmp_path: Path) -> Path:
    """Write a synthetic 3-row roster CSV matching data/schools/pssp_psrp_50.csv's columns."""
    path = tmp_path / "mini_roster.csv"
    fieldnames = [
        "sr_no", "emis_code", "school_name", "project", "cluster",
        "tehsil", "district", "students", "lat", "lon", "coords_raw",
    ]
    rows = [
        {
            "sr_no": 1, "emis_code": 90000001, "school_name": "Test Primary A",
            "project": "PSSP", "cluster": "C-1", "tehsil": "Test Tehsil",
            "district": "Test District", "students": 100, "lat": 32.08,
            "lon": 74.15, "coords_raw": "32.08N 74.15E",
        },
        {
            "sr_no": 2, "emis_code": 90000002, "school_name": "Test Primary B",
            "project": "PSSP", "cluster": "C-1", "tehsil": "Test Tehsil",
            "district": "Test District", "students": 200, "lat": 32.10,
            "lon": 74.18, "coords_raw": "32.10N 74.18E",
        },
        {
            "sr_no": 3, "emis_code": 90000003, "school_name": "Test Primary C",
            "project": "PSRP", "cluster": "C-2", "tehsil": "Test Tehsil",
            "district": "Test District", "students": 50, "lat": 32.20,
            "lon": 74.30, "coords_raw": "32.20N 74.30E",
        },
    ]
    with path.open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(rows)
    return path


def _write_year_sheet(ws, emis_col_header: str, months: dict[str, list[float]]) -> None:
    """Helper: write a SAR-shaped sheet.

    Row 1: spacer/title (mimics real workbook's extra header noise).
    Row 2: header row with an 'EMIS' cell + one datetime cell per month.
    Rows 3+: one row per school with attendance fractions per month.
    """
    ws.cell(row=1, column=1, value="Premier DLC — School Attendance Report (synthetic fixture)")

    month_names = list(months.keys())
    ws.cell(row=2, column=1, value=emis_col_header)
    ws.cell(row=2, column=2, value="School Name")
    for i, _name in enumerate(month_names):
        month_num = i + 1
        cell = ws.cell(row=2, column=3 + i, value=dt.datetime(2023, month_num, 1))
        cell.number_format = "mmm-yy"

    n_schools = len(next(iter(months.values())))
    for r in range(n_schools):
        row_idx = 3 + r
        emis = 90000001 + r
        ws.cell(row=row_idx, column=1, value=emis)
        ws.cell(row=row_idx, column=2, value=f"Test Primary {chr(65 + r)}")
        for i, name in enumerate(month_names):
            ws.cell(row=row_idx, column=3 + i, value=months[name][r])


@pytest.fixture
def mini_sar_xlsx(tmp_path: Path) -> Path:
    """Build a tiny, synthetic 3-year SAR workbook (openpyxl) for attendance parsing tests.

    Layout deliberately mirrors real-world quirks the parser is designed to survive:
      - an extra title row before the header
      - datetime month-header cells rather than fixed column offsets
      - school C is present in 2023/2024 but MISSING from 2025 (exercises the
        missing-school / null-monthly code path)
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws23 = wb.create_sheet("2023")
    _write_year_sheet(ws23, "EMIS", {
        "Jan": [0.90, 0.80, 0.70],
        "Feb": [0.92, 0.82, 0.72],
    })

    ws24 = wb.create_sheet("2024")
    _write_year_sheet(ws24, "EMIS", {
        "Jan": [0.88, 0.78, 0.68],
        "Feb": [0.91, 0.81, 0.71],
    })

    ws25 = wb.create_sheet("2025")
    # Only 2 schools this year (school C dropped) — exercises missing-school handling.
    ws25.cell(row=1, column=1, value="Premier DLC — School Attendance Report (synthetic fixture)")
    ws25.cell(row=2, column=1, value="EMIS")
    ws25.cell(row=2, column=2, value="School Name")
    ws25.cell(row=2, column=3, value=dt.datetime(2023, 1, 1)).number_format = "mmm-yy"
    ws25.cell(row=2, column=4, value=dt.datetime(2023, 2, 1)).number_format = "mmm-yy"
    ws25.cell(row=3, column=1, value=90000001)
    ws25.cell(row=3, column=2, value="Test Primary A")
    ws25.cell(row=3, column=3, value=0.95)
    ws25.cell(row=3, column=4, value=0.90)
    ws25.cell(row=4, column=1, value=90000002)
    ws25.cell(row=4, column=2, value="Test Primary B")
    ws25.cell(row=4, column=3, value=0.60)
    ws25.cell(row=4, column=4, value=0.65)

    path = tmp_path / "mini_sar.xlsx"
    wb.save(path)
    return path
